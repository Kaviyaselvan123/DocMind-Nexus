import os
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional
import pymupdf as fitz
import docx
import logging

logger = logging.getLogger(__name__)

CATEGORIES = {
    "Contract": ["agreement", "contract", "parties", "indemnification", "jurisdiction", "terms and conditions", "nda", "confidentiality"],
    "Financial": ["revenue", "ebitda", "balance sheet", "fiscal", "audit", "financial statement", "cash flow", "profit", "loss"],
    "Invoice": ["invoice", "bill to", "due date", "amount due", "remittance", "tax id", "subtotal", "purchase order", "po number"],
    "Policy": ["policy", "guidelines", "compliance", "code of conduct", "standard operating procedure", "sop", "regulatory", "governance"],
    "Technical": ["architecture", "api", "endpoint", "infrastructure", "kubernetes", "database", "specification", "sdk", "latency"],
    "HR": ["employment", "employee", "offer letter", "benefits", "salary", "pto", "severance", "job title", "probation"]
}

def extract_text_from_file(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    text = ""

    if suffix == ".pdf":
        try:
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"
        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {e}")
    elif suffix in [".docx", ".doc"]:
        try:
            doc = docx.Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text])
            # also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += "\n" + row_text
        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path}: {e}")
    else:
        # Plain text, markdown, csv, etc.
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            logger.error(f"Error reading text file {file_path}: {e}")

    return text.strip()

def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """Splits text into chunks of ~chunk_size words with overlap."""
    if not text:
        return []
    words = text.split()
    if len(words) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    step = max(1, chunk_size - overlap)
    while start < len(words):
        chunk_words = words[start:start + chunk_size]
        chunks.append(" ".join(chunk_words))
        start += step
    return chunks

def classify_text(text: str) -> Tuple[str, float]:
    """Classifies text into categories and assigns a confidence score."""
    if not text:
        return "General", 0.50

    lower_text = text.lower()
    scores = {}
    total_matches = 0
    
    for category, keywords in CATEGORIES.items():
        count = sum(len(re.findall(r"\b" + re.escape(kw) + r"\b", lower_text)) for kw in keywords)
        scores[category] = count
        total_matches += count

    if total_matches == 0:
        return "General", 0.50

    top_cat = max(scores, key=scores.get)
    best_score = scores[top_cat]
    confidence = min(0.98, max(0.55, best_score / (total_matches + 1) + 0.35))
    return top_cat, round(confidence, 2)

def extract_metadata_heuristics(text: str) -> Dict[str, Any]:
    """Extracts expiration dates, signature presence, and relevant tags."""
    has_signature = False
    sig_patterns = [
        r"(signed\s+by|signature\s*:|authorized\s+signatory|duly\s+authorized|/s/|accepted\s+and\s+agreed)",
        r"(signature\s+of|signee|by\s*:\s*[A-Z][a-z]+)"
    ]
    for pattern in sig_patterns:
        if re.search(pattern, text, re.IGNORECASE):
            has_signature = True
            break

    # Look for dates near expiration keywords
    expiration_date = None
    exp_match = re.search(
        r"(?:expires|expiration|valid\s+until|expiry|termination\s+date|term\s+ends)[\s:]*([A-Za-z]+ \d{1,2}, \d{4}|\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{4})",
        text,
        re.IGNORECASE
    )
    if exp_match:
        date_str = exp_match.group(1).strip()
        for fmt in ("%B %d, %Y", "%b %d, %Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                expiration_date = datetime.strptime(date_str, fmt)
                break
            except ValueError:
                continue

    # Tags extraction
    found_tags = set()
    tag_keywords = {
        "vendor": "Vendor",
        "nda": "NDA",
        "confidential": "Confidential",
        "sla": "SLA",
        "cloud": "Cloud",
        "security": "Security",
        "compliance": "Compliance",
        "finance": "Finance",
        "invoice": "Invoice",
        "annual": "Annual",
        "quarterly": "Quarterly",
        "hr": "HR"
    }
    lower = text.lower()
    for kw, tag in tag_keywords.items():
        if re.search(r"\b" + re.escape(kw) + r"\b", lower):
            found_tags.add(tag)

    return {
        "has_signature": has_signature,
        "expiration_date": expiration_date,
        "tags": list(found_tags)
    }
